# -*- coding: utf-8 -*-
"""病例管理路由 /cpx/case/*"""

import io
from datetime import datetime
from typing import Annotated
from urllib.parse import quote

from fastapi import APIRouter, Body, Depends, Path, Query, Request
from fastapi.responses import JSONResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

from app.common.response import SuccessResponse
from app.core.dependencies import db_getter

from app.plugin.module_cpx.auth.dependencies import ROLE_ADMIN, BizAuth, BusinessRole
from app.plugin.module_cpx.case.schema import CaseCreateSchema, CaseUpdateSchema
from app.plugin.module_cpx.case.service import CaseService
from app.plugin.module_cpx.fields import FIELDS
from app.plugin.module_cpx.log.service import LogService, get_client_ip

CaseRouter = APIRouter(prefix="/case", tags=["病例管理"])

_STATUS_TEXT = {"draft": "草稿", "submitted": "待审核", "approved": "审核通过", "rejected": "审核驳回"}


def _build_case_excel(items: list[dict]) -> Workbook:
    """把病例数据（基础字段 + 填报 form_data）生成 Excel 工作簿"""
    base_cols = [
        ("case_no", "病例编号"), ("patient_name", "患者姓名"), ("gender", "性别"), ("age", "年龄"),
        ("phone", "联系电话"), ("come_type", "来院方式"), ("diagnose_type", "诊断类型"),
        ("hospital_name", "所属医院"), ("doctor_name", "提交医生"), ("status", "审核状态"),
        ("create_time", "创建时间"),
    ]
    # 填报字段列：按标准字段字典顺序（fields.py），值取 form_data[field_code]；已出现在基础列的字段不重复导出
    base_codes = {code for code, _ in base_cols}
    form_cols = [(f["code"], f["name"]) for f in FIELDS if f["code"] not in base_codes]

    wb = Workbook()
    ws = wb.active
    ws.title = "病例数据"
    header = [name for _, name in base_cols] + [name for _, name in form_cols]
    ws.append(header)

    # 表头样式
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2563EB")
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for it in items:
        row = []
        for code, _ in base_cols:
            val = it.get(code)
            if code == "status":
                val = _STATUS_TEXT.get(val, val or "")
            row.append("" if val is None else val)
        fd = it.get("form_data") or {}
        for code, _ in form_cols:
            v = fd.get(code)
            row.append("" if v is None else v)
        ws.append(row)

    # 列宽自适应（表头与数据取最大，限 8~40）
    from openpyxl.utils import get_column_letter

    for col_idx, h in enumerate(header, start=1):
        max_len = len(str(h))
        for row_cells in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for c in row_cells:
                if c.value is not None:
                    max_len = max(max_len, len(str(c.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(8, max_len + 2), 40)
    ws.freeze_panes = "A2"
    return wb


def _build_case_pdf(items: list[dict], buf: io.BytesIO) -> None:
    """把病例数据生成多页横版 PDF（基础信息 + 按列分块的填报字段）。

    全字段（11 基础 + 72 填报 ≈ 83 列）远超单页宽度，因此：
    - 基础信息单独成表（含全部标识列）；
    - 填报字段按列分块，每块前 2 列重复「病例编号 / 患者姓名」便于跨页关联；
    - 每个表在纵向自动分页，表头随页重复。
    """
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak)

    # 中文支持：注册 Adobe 内置 CID 字体，无需外部字体文件
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    FONT = "STSong-Light"

    base_cols = [
        ("case_no", "病例编号"), ("patient_name", "患者姓名"), ("gender", "性别"), ("age", "年龄"),
        ("phone", "联系电话"), ("come_type", "来院方式"), ("diagnose_type", "诊断类型"),
        ("hospital_name", "所属医院"), ("doctor_name", "提交医生"), ("status", "审核状态"),
        ("create_time", "创建时间"),
    ]
    base_codes = {c for c, _ in base_cols}
    form_cols = [(f["code"], f["name"]) for f in FIELDS if f["code"] not in base_codes]
    # 填报字段分块时，每块前置的标识列（与基础表重复，便于跨页关联）
    identity = [("case_no", "病例编号"), ("patient_name", "患者姓名")]
    id_codes = {c for c, _ in identity}

    def text(it: dict, code: str) -> str:
        if code == "status":
            return _STATUS_TEXT.get(it.get(code), it.get(code) or "")
        if code in base_codes or code in id_codes:
            v = it.get(code)
        else:
            v = (it.get("form_data") or {}).get(code)
        return "" if v is None else str(v)

    head_style = ParagraphStyle("th", fontName=FONT, fontSize=7, leading=9, textColor=colors.white, alignment=1)
    body_style = ParagraphStyle("td", fontName=FONT, fontSize=7, leading=9, textColor=colors.black, alignment=0)
    title_style = ParagraphStyle("ti", fontName=FONT, fontSize=11, leading=14, textColor=colors.black, alignment=0)

    PAGE_W, _ = landscape(A4)
    MARGIN = 16
    usable = PAGE_W - 2 * MARGIN
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=MARGIN, rightMargin=MARGIN,
        topMargin=MARGIN, bottomMargin=MARGIN, title="病例数据导出",
    )

    def make_section(col_specs, title):
        n = len(col_specs)
        widths = [95.0 if i < 2 else (usable - 2 * 95.0) / (n - 2) for i in range(n)]
        data = [[Paragraph(name, head_style) for _, name in col_specs]]
        for it in items:
            data.append([Paragraph(text(it, code), body_style) for code, _ in col_specs])
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9CA3AF")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 1.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
        ]))
        return [Paragraph(title, title_style), Spacer(1, 4), t]

    flowables = []
    if not items:
        flowables.append(Paragraph("未查询到符合条件的病例数据", title_style))
    else:
        flowables += make_section(base_cols, "一、病例基础信息")
        flowables.append(PageBreak())
        chunk = 12
        total = (len(form_cols) + chunk - 1) // chunk
        for idx in range(0, len(form_cols), chunk):
            part = form_cols[idx:idx + chunk]
            specs = identity + part
            sec_title = f"二、填报字段（第 {idx // chunk + 1} / {total} 组，含标识列）"
            flowables += make_section(specs, sec_title)
            flowables.append(PageBreak())

    doc.build(flowables)


def _build_case_detail_pdf(detail: dict, buf: io.BytesIO) -> None:
    """生成单个病例的可读报告 PDF（A4 竖版，适合打印/归档）。

    结构：标题 → ① 基础信息 → ② 患者信息 → ③ 救治过程(填报数据) → ④ 审核信息。
    字段展示口径与 Web 详情弹窗保持一致（field_dict 优先，回退 template_fields；
    住院号/出院日期已在「患者信息」展示故不重复；select 字段反查 label）。
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable)

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    FONT = "STSong-Light"

    title_style = ParagraphStyle("title", fontName=FONT, fontSize=15, leading=20, alignment=1, textColor=colors.HexColor("#1E3A8A"))
    sub_style = ParagraphStyle("sub", fontName=FONT, fontSize=9, leading=12, alignment=1, textColor=colors.HexColor("#6B7280"))
    sec_style = ParagraphStyle("sec", fontName=FONT, fontSize=11, leading=15, textColor=colors.HexColor("#2563EB"), spaceBefore=8, spaceAfter=4)
    kv_key = ParagraphStyle("kvk", fontName=FONT, fontSize=8.5, leading=12, textColor=colors.HexColor("#6B7280"))
    kv_val = ParagraphStyle("kvv", fontName=FONT, fontSize=9, leading=12, textColor=colors.black)
    cell_style = ParagraphStyle("cell", fontName=FONT, fontSize=8.5, leading=12, textColor=colors.black)
    head_cell = ParagraphStyle("hc", fontName=FONT, fontSize=8.5, leading=12, textColor=colors.white, alignment=1)
    foot_style = ParagraphStyle("foot", fontName=FONT, fontSize=7.5, leading=10, textColor=colors.HexColor("#9CA3AF"), alignment=1)

    doc = SimpleDocTemplate(
        buf, pagesize=A4, topMargin=16 * mm, bottomMargin=14 * mm,
        leftMargin=16 * mm, rightMargin=16 * mm, title="病例报告",
    )
    usable = A4[0] - 32 * mm
    flow: list = []

    flow.append(Paragraph("胸痛中心 · 病例报告", title_style))
    flow.append(Paragraph(
        f"病例编号：{detail.get('case_no') or '-'}　|　状态：{_STATUS_TEXT.get(detail.get('status'), detail.get('status') or '-')}",
        sub_style,
    ))
    flow.append(Spacer(1, 6))
    flow.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#2563EB")))
    flow.append(Spacer(1, 6))

    def kv_table(pairs):
        data = []
        row = []
        for k, v in pairs:
            row.append(Paragraph(k, kv_key))
            row.append(Paragraph("-" if v in (None, "") else str(v), kv_val))
            if len(row) == 4:
                data.append(row)
                row = []
        if row:
            while len(row) < 4:
                row.append(Paragraph("", kv_key))
            data.append(row)
        t = Table(data, colWidths=[usable * 0.16, usable * 0.34, usable * 0.16, usable * 0.34])
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F3F4F6")),
            ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F3F4F6")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    # ① 基础信息
    flow.append(Paragraph("一、病例基础信息", sec_style))
    flow.append(kv_table([
        ("所属医院", detail.get("hospital_name")),
        ("提交医生", detail.get("doctor_name")),
        ("来院方式", detail.get("come_type")),
        ("诊断类型", detail.get("diagnose_type")),
        ("病例状态", _STATUS_TEXT.get(detail.get("status"), detail.get("status") or "-")),
        ("提交时间", detail.get("create_time")),
    ]))
    flow.append(Spacer(1, 6))

    # ② 患者信息
    flow.append(Paragraph("二、完整患者信息", sec_style))
    age = detail.get("age")
    age_s = f"{age} 岁" if age is not None else None
    fd = detail.get("form_data") or {}
    flow.append(kv_table([
        ("患者姓名", detail.get("patient_name")),
        ("性别", detail.get("gender")),
        ("年龄", age_s),
        ("联系电话", detail.get("phone")),
        ("住院号", fd.get("inpatient_no")),
        ("出院日期", fd.get("discharge_date")),
    ]))
    flow.append(Spacer(1, 6))

    # ③ 救治过程（填报数据）
    flow.append(Paragraph("三、救治过程（填报数据）", sec_style))
    dict_src = detail.get("field_dict") or detail.get("template_fields") or []
    seen: set = set()
    filled = []
    for f in dict_src:
        code = f.get("field_code")
        if not code or code in seen:
            continue
        seen.add(code)
        if code in ("inpatient_no", "discharge_date"):
            continue
        raw = fd.get(code)
        if raw in (None, ""):
            continue
        v = str(raw)
        if f.get("field_type") == "select":
            for o in (f.get("field_options") or []):
                if isinstance(o, dict):
                    if o.get("value") == v and o.get("label"):
                        v = o.get("label")
                        break
                elif o == v:
                    v = o
                    break
        filled.append((f.get("field_name") or code, v))

    if filled:
        table_data = [[Paragraph("字段", head_cell), Paragraph("内容", head_cell)]]
        for name, val in filled:
            table_data.append([Paragraph(name, kv_key), Paragraph(val, kv_val)])
        t = Table(table_data, colWidths=[usable * 0.30, usable * 0.70], repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        flow.append(t)
    else:
        flow.append(Paragraph("（该病例暂无填报数据）", foot_style))
    flow.append(Spacer(1, 6))

    # ④ 审核信息
    flow.append(Paragraph("四、审核信息", sec_style))
    audits = detail.get("audit_records") or []
    if audits:
        adata = [[Paragraph(h, head_cell) for h in ["审核人员", "所属医院", "审核结果", "审核意见", "审核时间"]]]
        for a in audits:
            res = a.get("audit_result")
            res_s = "通过" if res == "pass" else "驳回" if res == "reject" else (res or "-")
            adata.append([
                Paragraph(a.get("auditor_name") or "-", cell_style),
                Paragraph(a.get("auditor_hospital_name") or "-", cell_style),
                Paragraph(res_s, cell_style),
                Paragraph(a.get("audit_comment") or "-", cell_style),
                Paragraph(a.get("audit_time") or "-", cell_style),
            ])
        at = Table(adata, colWidths=[usable * 0.16, usable * 0.18, usable * 0.12, usable * 0.36, usable * 0.18], repeatRows=1)
        at.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        flow.append(at)
    else:
        flow.append(Paragraph("（该病例暂无审核记录）", foot_style))

    flow.append(Spacer(1, 8))
    flow.append(Paragraph(
        f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}　·　本文件由智慧胸痛中心管理系统自动生成",
        foot_style,
    ))
    doc.build(flow)


@CaseRouter.get("/list", summary="分页查询病例")
async def get_case_list_controller(
    auth: Annotated[BizAuth, Depends(BusinessRole([ROLE_ADMIN]))],
    db: Annotated[AsyncSession, Depends(db_getter)],
    page_no: int = Query(default=1, ge=1, description="页码"),
    page_size: int = Query(default=10, ge=1, le=100, description="每页条数"),
    hospital_id: int | None = Query(default=None, description="医院ID"),
    doctor_id: int | None = Query(default=None, description="医生ID"),
    case_no: str | None = Query(default=None, description="病例编号（完整精确/片段模糊）"),
    patient_name: str | None = Query(default=None, description="患者姓名（模糊）"),
    doctor_name: str | None = Query(default=None, description="医生姓名（模糊）"),
    status: str | None = Query(default=None, description="状态"),
    start_time: str | None = Query(default=None, description="开始时间 yyyy-MM-dd"),
    end_time: str | None = Query(default=None, description="结束时间 yyyy-MM-dd"),
) -> JSONResponse:
    service = CaseService(auth, db)
    result = await service.page(
        page_no=page_no,
        page_size=page_size,
        hospital_id=hospital_id,
        doctor_id=doctor_id,
        case_no=case_no,
        patient_name=patient_name,
        doctor_name=doctor_name,
        status=status,
        start_time=start_time,
        end_time=end_time,
    )
    return SuccessResponse(data=result, msg="查询病例列表成功")


@CaseRouter.get("/export", summary="导出病例数据 Excel")
async def export_case_controller(
    auth: Annotated[BizAuth, Depends(BusinessRole([ROLE_ADMIN]))],
    db: Annotated[AsyncSession, Depends(db_getter)],
    request: Request,
    hospital_id: int | None = Query(default=None, description="医院ID"),
    doctor_id: int | None = Query(default=None, description="医生ID"),
    case_no: str | None = Query(default=None, description="病例编号（完整精确/片段模糊）"),
    patient_name: str | None = Query(default=None, description="患者姓名（模糊）"),
    doctor_name: str | None = Query(default=None, description="医生姓名（模糊）"),
    status: str | None = Query(default=None, description="状态"),
    start_time: str | None = Query(default=None, description="开始时间 yyyy-MM-dd"),
    end_time: str | None = Query(default=None, description="结束时间 yyyy-MM-dd"),
) -> Response:
    service = CaseService(auth, db)
    items = await service.export_data(
        hospital_id=hospital_id,
        doctor_id=doctor_id,
        case_no=case_no,
        patient_name=patient_name,
        doctor_name=doctor_name,
        status=status,
        start_time=start_time,
        end_time=end_time,
    )
    await LogService.create(
        db,
        user_id=auth.user.id,
        module="病例管理",
        operation="导出Excel",
        description=f"用户 {auth.user.real_name} 导出病例数据（{len(items)} 条）",
        ip_address=get_client_ip(request),
    )
    wb = _build_case_excel(items)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"病例导出_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"},
    )


@CaseRouter.get("/export_pdf", summary="导出病例数据 PDF")
async def export_case_pdf_controller(
    auth: Annotated[BizAuth, Depends(BusinessRole([ROLE_ADMIN]))],
    db: Annotated[AsyncSession, Depends(db_getter)],
    request: Request,
    hospital_id: int | None = Query(default=None, description="医院ID"),
    doctor_id: int | None = Query(default=None, description="医生ID"),
    case_no: str | None = Query(default=None, description="病例编号（完整精确/片段模糊）"),
    patient_name: str | None = Query(default=None, description="患者姓名（模糊）"),
    doctor_name: str | None = Query(default=None, description="医生姓名（模糊）"),
    status: str | None = Query(default=None, description="状态"),
    start_time: str | None = Query(default=None, description="开始时间 yyyy-MM-dd"),
    end_time: str | None = Query(default=None, description="结束时间 yyyy-MM-dd"),
) -> Response:
    service = CaseService(auth, db)
    items = await service.export_data(
        hospital_id=hospital_id,
        doctor_id=doctor_id,
        case_no=case_no,
        patient_name=patient_name,
        doctor_name=doctor_name,
        status=status,
        start_time=start_time,
        end_time=end_time,
    )
    await LogService.create(
        db,
        user_id=auth.user.id,
        module="病例管理",
        operation="导出PDF",
        description=f"用户 {auth.user.real_name} 导出病例数据（{len(items)} 条）",
        ip_address=get_client_ip(request),
    )
    buf = io.BytesIO()
    _build_case_pdf(items, buf)
    buf.seek(0)
    fname = f"病例导出_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return Response(
        content=buf.read(),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"},
    )


@CaseRouter.get("/export_pdf/{id}", summary="导出单个病例报告 PDF")
async def export_case_pdf_detail_controller(
    auth: Annotated[BizAuth, Depends(BusinessRole([ROLE_ADMIN]))],
    id: Annotated[int, Path(description="病例ID")],
    db: Annotated[AsyncSession, Depends(db_getter)],
    request: Request,
) -> Response:
    service = CaseService(auth, db)
    detail = await service.detail(id=id)
    await LogService.create(
        db,
        user_id=auth.user.id,
        module="病例管理",
        operation="导出PDF(单病例)",
        description=f"用户 {auth.user.real_name} 导出病例 {detail.get('case_no')}（{detail.get('patient_name')}）报告",
        ip_address=get_client_ip(request),
    )
    buf = io.BytesIO()
    _build_case_detail_pdf(detail, buf)
    buf.seek(0)
    fname = f"病例报告_{detail.get('case_no') or id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return Response(
        content=buf.read(),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"},
    )


@CaseRouter.get("/detail/{id}", summary="获取病例详情")
async def get_case_detail_controller(
    auth: Annotated[BizAuth, Depends(BusinessRole([ROLE_ADMIN]))],
    id: Annotated[int, Path(description="病例ID")],
    db: Annotated[AsyncSession, Depends(db_getter)],
) -> JSONResponse:
    service = CaseService(auth, db)
    result = await service.detail(id=id)
    return SuccessResponse(data=result, msg="获取病例详情成功")


@CaseRouter.get("/timeline/{id}", summary="救治时间轴（含关键质控指标）")
async def get_case_timeline_controller(
    auth: Annotated[BizAuth, Depends(BusinessRole([ROLE_ADMIN]))],
    id: Annotated[int, Path(description="病例ID")],
    db: Annotated[AsyncSession, Depends(db_getter)],
) -> JSONResponse:
    service = CaseService(auth, db)
    result = await service.timeline(id=id)
    return SuccessResponse(data=result, msg="获取时间轴成功")


@CaseRouter.get("/analysis/{id}", summary="单病例分析（质控指标校验）")
async def get_case_analysis_controller(
    auth: Annotated[BizAuth, Depends(BusinessRole([ROLE_ADMIN]))],
    id: Annotated[int, Path(description="病例ID")],
    db: Annotated[AsyncSession, Depends(db_getter)],
) -> JSONResponse:
    service = CaseService(auth, db)
    result = await service.analysis(id=id)
    return SuccessResponse(data=result, msg="获取病例分析成功")


@CaseRouter.get("/followup/{id}", summary="病例随访记录详情")
async def get_case_followup_controller(
    auth: Annotated[BizAuth, Depends(BusinessRole([ROLE_ADMIN]))],
    id: Annotated[int, Path(description="病例ID")],
    db: Annotated[AsyncSession, Depends(db_getter)],
) -> JSONResponse:
    service = CaseService(auth, db)
    result = await service.followup(id=id)
    return SuccessResponse(data=result, msg="获取随访记录成功")


@CaseRouter.post("/create", summary="新增病例（测试/联调）")
async def create_case_controller(
    auth: Annotated[BizAuth, Depends(BusinessRole([ROLE_ADMIN]))],
    data: Annotated[CaseCreateSchema, Body(description="新增参数")],
    db: Annotated[AsyncSession, Depends(db_getter)],
    request: Request,
) -> JSONResponse:
    service = CaseService(auth, db)
    result = await service.create(data=data)
    await LogService.create(
        db,
        user_id=auth.user.id,
        module="病例管理",
        operation="新增病例",
        description=f"新增病例：{data.patient_name or '未命名'}",
        ip_address=get_client_ip(request),
    )
    return SuccessResponse(data=result, msg="新增病例成功")


@CaseRouter.post("/submit/{id}", summary="提交病例（草稿/驳回 → 待审核）")
async def submit_case_controller(
    auth: Annotated[BizAuth, Depends(BusinessRole([ROLE_ADMIN]))],
    id: Annotated[int, Path(description="病例ID")],
    db: Annotated[AsyncSession, Depends(db_getter)],
) -> JSONResponse:
    service = CaseService(auth, db)
    result = await service.submit(id=id)
    return SuccessResponse(data=result, msg="提交病例成功")


@CaseRouter.put("/{id}", summary="更新病例基础信息（web 与 APP 互通）")
async def update_case_controller(
    auth: Annotated[BizAuth, Depends(BusinessRole([ROLE_ADMIN]))],
    id: Annotated[int, Path(description="病例ID")],
    data: Annotated[CaseUpdateSchema, Body(description="更新参数（仅传入字段生效）")],
    db: Annotated[AsyncSession, Depends(db_getter)],
    request: Request,
) -> JSONResponse:
    service = CaseService(auth, db)
    result = await service.update(id=id, data=data)
    await LogService.create(
        db,
        user_id=auth.user.id,
        module="病例管理",
        operation="更新病例基础信息",
        description=f"更新病例 {id}：电话={data.phone or '-'}",
        ip_address=get_client_ip(request),
    )
    return SuccessResponse(data=result, msg="更新病例成功")


@CaseRouter.delete("/{id}", summary="删除病例（级联删除详情/审核/随访/心电）")
async def delete_case_controller(
    auth: Annotated[BizAuth, Depends(BusinessRole([ROLE_ADMIN]))],
    id: Annotated[int, Path(description="病例ID")],
    db: Annotated[AsyncSession, Depends(db_getter)],
    request: Request,
) -> JSONResponse:
    service = CaseService(auth, db)
    result = await service.delete(id=id)
    await LogService.create(
        db,
        user_id=auth.user.id,
        module="病例管理",
        operation="删除病例",
        description=f"删除病例 {result['case_no']}（{result['patient_name']}），连带 {result['deleted']} 条关联记录",
        ip_address=get_client_ip(request),
    )
    return SuccessResponse(data=result, msg="病例已删除")
